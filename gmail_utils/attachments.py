import base64
import logging
from io import BytesIO
from typing import List, Dict, Tuple

from googleapiclient.errors import HttpError
from config.settings import ATTACHMENT_SETTINGS

# Optional imports for parsing attachments
try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except Exception:
    pdf_extract_text = None

try:
    import docx  # python-docx
except Exception:
    docx = None

try:
    import openpyxl
except Exception:
    openpyxl = None

logger = logging.getLogger(__name__)


def _iter_parts(payload: Dict) -> List[Dict]:
    """Recursively collect all parts from a Gmail message payload."""
    parts = []
    if not payload:
        return parts
    if "parts" in payload and isinstance(payload["parts"], list):
        for p in payload["parts"]:
            parts.append(p)
            parts.extend(_iter_parts(p))
    return parts


def find_attachments(payload: Dict) -> List[Dict]:
    """Find attachment parts in the Gmail message payload."""
    attachments = []
    for part in [payload] + _iter_parts(payload):
        filename = part.get("filename", "")
        body = part.get("body", {})
        mime = part.get("mimeType", "")
        # Attachment parts usually have a filename and an attachmentId
        attachment_id = body.get("attachmentId")
        size = body.get("size", 0)
        if filename and attachment_id:
            attachments.append({
                "filename": filename,
                "mimeType": mime,
                "attachmentId": attachment_id,
                "size": size,
            })
    return attachments


def download_attachment(service, message_id: str, attachment_id: str) -> bytes:
    """Download attachment bytes using Gmail API."""
    att = service.users().messages().attachments().get(
        userId="me", messageId=message_id, id=attachment_id
    ).execute()
    data = att.get("data")
    if not data:
        return b""
    try:
        return base64.urlsafe_b64decode(data)
    except Exception:
        # Some payloads may use standard base64
        return base64.b64decode(data)


def extract_text_from_attachment(data: bytes, mime_type: str, filename: str) -> str:
    """Extract text from attachment based on MIME type."""
    max_len = ATTACHMENT_SETTINGS["MAX_TEXT_LENGTH"]

    try:
        if mime_type == "application/pdf" and pdf_extract_text:
            text = pdf_extract_text(BytesIO(data))
            return (text or "").strip()[:max_len]

        if mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ) and docx:
            doc = docx.Document(BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            return text.strip()[:max_len]

        if mime_type in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        ) and openpyxl:
            wb = openpyxl.load_workbook(filename=BytesIO(data), read_only=True, data_only=True)
            text_chunks = []
            for ws in wb.worksheets[:2]:  # limit sheets
                rows_iter = ws.iter_rows(min_row=1, max_row=50, values_only=True)
                for row in rows_iter:
                    row_vals = [str(v) if v is not None else "" for v in row]
                    if any(row_vals):
                        text_chunks.append("\t".join(row_vals))
            text = "\n".join(text_chunks)
            return text.strip()[:max_len]

        # Fallback: unsupported types return empty string
        return ""
    except Exception as e:
        logger.warning(f"Error extracting text from attachment {filename}: {e}")
        return ""


def process_message_attachments(service, message: Dict) -> Tuple[str, List[Dict]]:
    """Download and parse attachments for a Gmail message. Returns (text, metadata)."""
    payload = message.get("payload", {})
    att_settings = ATTACHMENT_SETTINGS
    supported = set(att_settings["SUPPORTED_MIME_TYPES"])
    max_size = att_settings["MAX_SIZE_BYTES"]

    attachments = find_attachments(payload)
    if not attachments:
        return "", []

    text_parts = []
    meta: List[Dict] = []

    for att in attachments:
        mime = att.get("mimeType", "")
        size = int(att.get("size", 0) or 0)
        filename = att.get("filename", "")
        attachment_id = att.get("attachmentId")

        meta.append({
            "filename": filename,
            "mimeType": mime,
            "size": size,
        })

        # Skip unsupported or oversized attachments
        if mime not in supported or size > max_size:
            logger.info(f"Skipping attachment {filename} (type={mime}, size={size})")
            continue

        try:
            data = download_attachment(service, message.get("id"), attachment_id)
            extracted = extract_text_from_attachment(data, mime, filename)
            if extracted:
                text_parts.append(f"Attachment {filename} ({mime}):\n{extracted}")
        except HttpError as e:
            logger.warning(f"Gmail API error downloading attachment {filename}: {e}")
        except Exception as e:
            logger.warning(f"Unexpected error processing attachment {filename}: {e}")

    combined_text = "\n\n".join(text_parts)
    return combined_text, meta